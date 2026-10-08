from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .config import Config
from .utils import separador


class GeneradorResumen2:
    """Responsable de generar la hoja RESUMEN2 con los indicadores del proceso."""

    def __init__(self, config: Config):
        self.config = config

    def generar(self, df_pqrsd, wb):
        """Genera la hoja RESUMEN2 con los indicadores del proceso."""
        separador("GENERANDO RESUMEN2")

        nombre_hoja = "RESUMEN2"

        if nombre_hoja in wb.sheetnames:
            del wb[nombre_hoja]

        if "RESUMEN" in wb.sheetnames:
            del wb["RESUMEN"]

        ws = wb.create_sheet(nombre_hoja)

        ws["A1"] = "RESUMEN PQRSD - PLANEACIÓN"
        ws["A1"].font = Font(
            bold=True,
            size=14,
        )

        ws["A3"] = "TOTAL DE REGISTROS"
        ws["B3"] = len(df_pqrsd)

        fila = 5

        conteo_estado = df_pqrsd[self.config.COL_ESTADO].value_counts(dropna=False)
        fila = self._escribir_conteo(
            ws,
            fila,
            "RESUMEN POR ESTADO",
            conteo_estado,
        )

        fila += 2

        conteo_area = df_pqrsd[self.config.COL_AREA_SALIDA].value_counts(dropna=False)
        fila = self._escribir_conteo(
            ws,
            fila,
            "RESUMEN POR ÁREA / OFICINA",
            conteo_area,
        )

        fila += 2

        ws.cell(fila, 1, "FILTRO APLICADO").font = Font(bold=True)

        fila += 1
        ws.cell(
            fila,
            1,
            "AREA RESPONSABLE o SECRETARIA contiene PLANEACION",
        )

        fila += 1
        ws.cell(
            fila,
            1,
            f"FECHA > {self.config.FECHA_CORTE.strftime('%d/%m/%Y')}",
        )

        fila += 2

        ws.cell(fila, 1, "FECHA DE GENERACIÓN")
        ws.cell(fila, 2, datetime.now(ZoneInfo("America/Bogota")).replace(tzinfo=None))
        ws.cell(
            fila,
            2,
        ).number_format = "dd/mm/yyyy hh:mm"

        self._configurar_ancho_resumen(ws)

        return wb

    def _escribir_conteo(self, ws, fila, titulo, serie):
        """Escribe una sección de conteo en una hoja de Excel."""
        ws.cell(fila, 1, titulo).font = Font(bold=True)
        fila += 1

        for valor, cantidad in serie.items():
            ws.cell(fila, 1, str(valor))
            ws.cell(fila, 2, int(cantidad))
            fila += 1

        return fila

    def _configurar_ancho_resumen(self, ws, ancho=40):
        """Configura el ancho de las columnas del resumen."""
        for columna in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(columna)].width = ancho


class FormateadorHojaPrincipal:
    """Responsable de aplicar formatos a la hoja principal."""

    def __init__(self, config: Config):
        self.config = config

    def aplicar(self, ws):
        """Aplica todos los formatos de la hoja principal."""
        self._formatear_encabezados(ws)
        self._formatear_fechas(ws)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 45

        self._calcular_ancho_columnas(ws)
        self._envolver_texto(ws)

    def _formatear_encabezados(self, ws):
        """Aplica formato a la fila de encabezados."""
        for celda in ws[1]:
            celda.font = Font(bold=True)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    def _formatear_fechas(self, ws):
        """Aplica formato dd/mm/yyyy a columnas de fecha."""
        columnas_fecha = [
            self.config.COL_FECHA,
            self.config.COL_FECHA_LIMITE,
            self.config.COL_FECHA_RESPUESTA,
            self.config.COL_FECHA_ENTREGA_RESPUESTA,
        ]

        encabezados = {
            ws.cell(1, columna).value: columna
            for columna in range(1, ws.max_column + 1)
        }

        for nombre_columna in columnas_fecha:
            numero_columna = encabezados.get(nombre_columna)

            if numero_columna is None:
                continue

            for fila in range(2, ws.max_row + 1):
                celda = ws.cell(fila, numero_columna)

                if celda.value is not None:
                    celda.number_format = "dd/mm/yyyy"

    def _calcular_ancho_columnas(self, ws):
        """Calcula el ancho de las columnas usando hasta 100 filas."""
        for columna in range(1, ws.max_column + 1):
            letra = get_column_letter(columna)
            max_length = 0

            for fila in range(1, min(ws.max_row, 100) + 1):
                valor = ws.cell(fila, columna).value

                if valor is not None:
                    max_length = max(
                        max_length,
                        len(str(valor)),
                    )

            ancho = min(
                max(max_length + 2, 12),
                45,
            )

            ws.column_dimensions[letra].width = ancho

    def _envolver_texto(self, ws):
        """Activa ajuste de texto y alineación vertical."""
        for fila in ws.iter_rows():
            for celda in fila:
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


class FormateadorResumen2:
    """Responsable de aplicar formato a la hoja RESUMEN2."""

    def __init__(self, config: Config):
        self.config = config

    def aplicar(self, wb):
        """Aplica ajuste de texto a la hoja RESUMEN2."""
        if "RESUMEN2" not in wb.sheetnames:
            return

        ws_resumen = wb["RESUMEN2"]

        for fila in ws_resumen.iter_rows():
            for celda in fila:
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )
